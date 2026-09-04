import os
import re
import pdfplumber
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import load_workbook
import subprocess
import ttkbootstrap as tb
from ttkbootstrap.constants import *
from pathlib import Path

# ======================
# Configuración global (LOCAL)
# ======================
APP_OWNER = "Angel"
BASE_DIR = Path(__file__).resolve().parent
PLANTILLAS_DIR = BASE_DIR / "plantillas"
if not PLANTILLAS_DIR.exists():
    PLANTILLAS_DIR = BASE_DIR  # fallback

EXCEL_REPORTE = str((PLANTILLAS_DIR / "REPORTE DE ANALISIS.xlsx").resolve())
EXCEL_SOCIEDADES = str((BASE_DIR / "SOCIEDADES.xlsx").resolve())

# LibreOffice (Mac)
SOFFICE_PATH = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

sociedades = []
datos_seleccionados = {}
entrega_folder = ""
procesadas = set()

# ======================
# Utilidades
# ======================
def _check_base_files():
    faltan = []
    if not os.path.exists(EXCEL_REPORTE):
        faltan.append(f"No encuentro REPORTE DE ANALISIS.xlsx en:\n{EXCEL_REPORTE}")
    if not os.path.exists(EXCEL_SOCIEDADES):
        faltan.append(f"No encuentro SOCIEDADES.xlsx en:\n{EXCEL_SOCIEDADES}")
    if not os.path.exists(SOFFICE_PATH):
        faltan.append(f"No encuentro LibreOffice (soffice) en:\n{SOFFICE_PATH}")
    if faltan:
        messagebox.showerror("Faltan archivos", "\n\n".join(faltan))
        return False
    return True

def _safe_float(s: str) -> float:
    value = str(s).strip()
    if not value:
        raise ValueError("valor vacio")
    value = re.sub(r"[^0-9,\.\-]", "", value)
    if value in ("", "-", ",", "."):
        raise ValueError("valor invalido")
    if "," in value and "." in value:
        if value.rfind(",") > value.rfind("."):
            value = value.replace(".", "").replace(",", ".")
        else:
            value = value.replace(",", "")
    else:
        value = value.replace(",", ".")
    return float(value)

def _convert_xlsx_to_pdf(outdir: str) -> str:
    """
    Convierte EXCEL_REPORTE a PDF con LibreOffice.
    Devuelve la ruta del PDF generado (detectado por nombre base).
    """
    os.makedirs(outdir, exist_ok=True)

    # Ejecutar LibreOffice sin shell para evitar problemas con espacios
    args = [
        SOFFICE_PATH,
        "--headless",
        "--convert-to", "pdf",
        "--outdir", outdir,
        EXCEL_REPORTE
    ]

    try:
        subprocess.run(args, check=True, capture_output=True, text=True)
    except subprocess.CalledProcessError as e:
        raise RuntimeError(
            "LibreOffice falló al convertir.\n\n"
            f"STDOUT:\n{e.stdout}\n\nSTDERR:\n{e.stderr}"
        )

    # LibreOffice genera el PDF con el nombre del xlsx base
    base_pdf = Path(EXCEL_REPORTE).with_suffix(".pdf").name  # "REPORTE DE ANALISIS.pdf"
    candidate = os.path.join(outdir, base_pdf)

    # Si no está exacto, buscamos cualquier PDF que empiece con el base (por si LO agrega sufijos)
    if os.path.exists(candidate):
        return candidate

    base_stem = Path(EXCEL_REPORTE).stem  # "REPORTE DE ANALISIS"
    pdfs = [p for p in os.listdir(outdir) if p.lower().endswith(".pdf") and p.startswith(base_stem)]
    if pdfs:
        return os.path.join(outdir, pdfs[0])

    raise FileNotFoundError("LibreOffice terminó pero no encontré el PDF generado en la carpeta de salida.")

def _limpiar_excel_reporte():
    wb = load_workbook(EXCEL_REPORTE)
    ws = wb.active
    for cell in ["I7","I9","I11","I13","B16","B21","A28","C28","E28","G28"]:
        ws[cell] = ""
    wb.save(EXCEL_REPORTE)
    wb.close()

# ======================
# Funciones principales
# ======================
def seleccionar_carpeta():
    global entrega_folder
    if not _check_base_files():
        return

    folder = filedialog.askdirectory(title="Selecciona la carpeta de ENTREGA")
    if not folder:
        return
    entrega_folder = folder
    detectar_preliminares()

def detectar_preliminares():
    global sociedades
    if not entrega_folder:
        return

    prelim_pdf = None
    for f in os.listdir(entrega_folder):
        if f.upper().startswith("PRELIMINARES") and f.lower().endswith(".pdf"):
            prelim_pdf = os.path.join(entrega_folder, f)
            break

    if not prelim_pdf:
        messagebox.showwarning("Atención", "No se encontró PRELIMINARES en la carpeta.")
        return

    with pdfplumber.open(prelim_pdf) as pdf:
        page = pdf.pages[0]
        raw = page.extract_text()
        if not raw:
            messagebox.showerror("Error", "No pude extraer texto del PRELIMINARES.pdf (extract_text vacío).")
            return
        text = raw.splitlines()

    sociedades = []
    for line in text:
        if "PROVEEDOR" in line or "BARRA" in line:
            continue
        partes = line.split()
        barra_idx = next((i for i, p in enumerate(partes) if "-" in p), None)
        if barra_idx is not None and barra_idx + 1 < len(partes):
            sociedad = " ".join(partes[:barra_idx])
            barra = partes[barra_idx]
            peso_inicial = partes[barra_idx + 1]
            peso_final = partes[-1]
            sociedades.append({
                "sociedad": sociedad,
                "barra": barra,
                "peso_inicial": peso_inicial,
                "peso_final": peso_final
            })

    combo_sociedades["values"] = [s['sociedad'] for s in sociedades]
    combo_sociedades.set("")
    tabla_procesadas.delete(*tabla_procesadas.get_children())

def seleccionar_sociedad(event=None):
    global datos_seleccionados
    idx = combo_sociedades.current()
    if idx < 0:
        return

    sociedad = sociedades[idx]["sociedad"]
    if sociedad in procesadas:
        messagebox.showwarning("Atención", f"{sociedad} ya fue procesada.")
        return

    datos_seleccionados = sociedades[idx]

    # 1) Leer SOCIEDADES.xlsx (LOCAL)
    wb_soc = load_workbook(EXCEL_SOCIEDADES, data_only=True)
    ws_soc = wb_soc.active
    nit, rucom, municipio = "", "", ""
    for row in ws_soc.iter_rows(min_row=2, values_only=True):
        if row and str(row[0]).strip() == sociedad:
            nit = str(row[1]) if row[1] else ""
            rucom = str(int(row[2])) if isinstance(row[2], (int, float)) else (str(row[2]) if row[2] else "")
            municipio = str(row[3]) if row[3] else ""
            break
    wb_soc.close()

    # 2) Llenar REPORTE DE ANALISIS.xlsx (plantilla LOCAL)
    wb = load_workbook(EXCEL_REPORTE)
    ws = wb.active
    ws["I7"], ws["I9"], ws["I11"], ws["I13"] = sociedad, nit, rucom, municipio
    ws["B16"], ws["B21"] = datos_seleccionados["barra"], datos_seleccionados["barra"]
    ws["A28"] = _safe_float(datos_seleccionados["peso_inicial"])
    ws["C28"] = _safe_float(datos_seleccionados["peso_final"])
    wb.save(EXCEL_REPORTE)
    wb.close()

def guardar_excel():
    if not datos_seleccionados:
        messagebox.showwarning("Atención", "Selecciona una sociedad primero.")
        return

    try:
        ley_au = _safe_float(campo_ley_au.get())
        ley_ag = _safe_float(campo_ley_ag.get())
    except ValueError:
        messagebox.showerror("Error", "Ley AU y Ley AG deben ser numericas. Puedes usar punto o coma decimal.")
        return

    wb = load_workbook(EXCEL_REPORTE)
    ws = wb.active
    ws["E28"], ws["G28"] = ley_au, ley_ag
    wb.save(EXCEL_REPORTE)
    wb.close()
    messagebox.showinfo("Éxito", "Leyes guardadas.")

def generar_pdf():
    if not datos_seleccionados:
        messagebox.showwarning("Atención", "Selecciona una sociedad.")
        return

    if not entrega_folder:
        messagebox.showwarning("Atención", "Selecciona la carpeta de ENTREGA.")
        return

    leyes_folder = os.path.join(entrega_folder, "LEYES")
    os.makedirs(leyes_folder, exist_ok=True)

    try:
        # Convertir Excel a PDF
        pdf_generado = _convert_xlsx_to_pdf(leyes_folder)

        barra = datos_seleccionados["barra"]
        destino = os.path.join(leyes_folder, f"REPORTE LEYES - {barra}.pdf")

        # Si existe, reemplazar
        if os.path.exists(destino):
            try:
                os.remove(destino)
            except:
                pass

        os.rename(pdf_generado, destino)

        # Limpiar Excel
        _limpiar_excel_reporte()

        # Registrar como procesada
        procesadas.add(datos_seleccionados["sociedad"])
        tabla_procesadas.insert("", "end", values=(datos_seleccionados["sociedad"], "✅"))
        combo_sociedades.set("")
        campo_ley_au.delete(0, tk.END)
        campo_ley_ag.delete(0, tk.END)

        messagebox.showinfo("Éxito", f"PDF generado:\n{destino}")

    except Exception as e:
        messagebox.showerror("Error PDF", f"No se pudo generar el PDF.\n\nDetalle:\n{e}")

# ======================
# Interfaz gráfica
# ======================

BG = "#F4F7FA"
PANEL = "#FFFFFF"
INK = "#172033"
MUTED = "#667085"
BORDER = "#D8E0EA"
ACCENT = "#0E3662"

def configure_enterprise_style(window):
    window.configure(bg=BG)
    style = tb.Style()
    style.configure("App.TFrame", background=BG)
    style.configure("Card.TFrame", background=PANEL, relief="solid", borderwidth=1)
    style.configure("Header.TFrame", background=ACCENT)
    style.configure("Title.TLabel", background=ACCENT, foreground="white", font=("Calibri", 23, "bold"))
    style.configure("Subtitle.TLabel", background=ACCENT, foreground="#DDE7F2", font=("Calibri", 11))
    style.configure("Section.TLabel", background=PANEL, foreground=INK, font=("Calibri", 13, "bold"))
    style.configure("Muted.TLabel", background=PANEL, foreground=MUTED, font=("Calibri", 10))
    style.configure("Field.TLabel", background=PANEL, foreground=INK, font=("Calibri", 12, "bold"))
    style.configure("Treeview", rowheight=28, font=("Calibri", 11))
    style.configure("Treeview.Heading", font=("Calibri", 11, "bold"))

def make_header(parent, title, subtitle):
    header = ttk.Frame(parent, style="Header.TFrame", padding=(24, 18))
    header.pack(fill="x", pady=(0, 14))
    ttk.Label(header, text=title, style="Title.TLabel").pack(anchor="w")
    ttk.Label(header, text=subtitle, style="Subtitle.TLabel").pack(anchor="w", pady=(3, 0))
    return header

def animate_window(window, step=0):
    try:
        alpha = min(1.0, 0.88 + step * 0.025)
        window.attributes("-alpha", alpha)
        if alpha < 1.0:
            window.after(20, lambda: animate_window(window, step + 1))
    except Exception:
        pass

def wire_button_hover(button, normal, hover):
    def enter(_event=None):
        try:
            if str(button.cget("state")) != "disabled":
                button.configure(bootstyle=hover)
        except Exception:
            pass
    def leave(_event=None):
        try:
            button.configure(bootstyle=normal)
        except Exception:
            pass
    button.bind("<Enter>", enter)
    button.bind("<Leave>", leave)
    return button

root = tb.Window(themename="flatly")
root.title(f"Angel | Reporte de Análisis")
root.geometry("860x650")
root.minsize(820, 620)
configure_enterprise_style(root)

main = tb.Frame(root, style="App.TFrame", padding=18)
main.pack(fill="both", expand=True)
make_header(main, "REPORTE DE ANÁLISIS", "Sistema Angel | Registro de leyes AU/AG y generación de PDFs")

toolbar = tb.Frame(main, style="Card.TFrame", padding=18)
toolbar.pack(fill="x", pady=(0, 14))
toolbar.columnconfigure(1, weight=1)
btn_carpeta = tb.Button(toolbar, text="SELECCIONAR ENTREGA", bootstyle=PRIMARY, command=seleccionar_carpeta)
btn_carpeta.grid(row=0, column=0, padx=(0, 14), ipady=5, sticky="ew")
wire_button_hover(btn_carpeta, PRIMARY, SUCCESS)
tb.Label(toolbar, text="Carga el PRELIMINARES de la entrega seleccionada", style="Muted.TLabel").grid(row=0, column=1, sticky="w")

frame_datos = tb.Frame(main, style="Card.TFrame", padding=20)
frame_datos.pack(fill="x", pady=(0, 14))
frame_datos.columnconfigure(1, weight=1)
tb.Label(frame_datos, text="Datos de entrada", style="Section.TLabel").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 12))

tb.Label(frame_datos, text="Sociedad", style="Field.TLabel").grid(row=1, column=0, sticky="e", padx=(0, 12), pady=8)
combo_sociedades = tb.Combobox(frame_datos, state="readonly", width=52)
combo_sociedades.grid(row=1, column=1, columnspan=2, padx=5, pady=8, sticky="ew")
combo_sociedades.bind("<<ComboboxSelected>>", seleccionar_sociedad)

tb.Label(frame_datos, text="Ley AU", style="Field.TLabel").grid(row=2, column=0, sticky="e", padx=(0, 12), pady=8)
campo_ley_au = tb.Entry(frame_datos, width=22)
campo_ley_au.grid(row=2, column=1, sticky="w", padx=5, pady=8)
tb.Label(frame_datos, text="Ley AG", style="Field.TLabel").grid(row=3, column=0, sticky="e", padx=(0, 12), pady=8)
campo_ley_ag = tb.Entry(frame_datos, width=22)
campo_ley_ag.grid(row=3, column=1, sticky="w", padx=5, pady=8)
tb.Label(frame_datos, text="Puedes usar 1.25 o 1,25", style="Muted.TLabel").grid(row=4, column=1, sticky="w", padx=5, pady=(2, 0))

frame_botones = tb.Frame(main, style="Card.TFrame", padding=18)
frame_botones.pack(fill="x", pady=(0, 14))
frame_botones.columnconfigure(0, weight=1)
frame_botones.columnconfigure(1, weight=1)
btn_guardar = tb.Button(frame_botones, text="GUARDAR LEYES", bootstyle=SUCCESS, command=guardar_excel)
btn_guardar.grid(row=0, column=0, padx=(0, 8), ipady=5, sticky="ew")
wire_button_hover(btn_guardar, SUCCESS, PRIMARY)
btn_pdf = tb.Button(frame_botones, text="GENERAR PDF", bootstyle=PRIMARY, command=generar_pdf)
btn_pdf.grid(row=0, column=1, padx=(8, 0), ipady=5, sticky="ew")
wire_button_hover(btn_pdf, PRIMARY, SUCCESS)

frame_lista = tb.Frame(main, style="Card.TFrame", padding=18)
frame_lista.pack(fill="both", expand=True)
tb.Label(frame_lista, text="Sociedades procesadas", style="Section.TLabel").pack(anchor="w", pady=(0, 10))
tabla_procesadas = ttk.Treeview(frame_lista, columns=("sociedad", "estado"), show="headings", height=8)
tabla_procesadas.heading("sociedad", text="Sociedad")
tabla_procesadas.heading("estado", text="Estado")
tabla_procesadas.column("sociedad", width=600)
tabla_procesadas.column("estado", width=120, anchor="center")
tabla_procesadas.pack(fill="both", expand=True)

animate_window(root)
root.mainloop()
