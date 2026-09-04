import os
import re
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pdfplumber
import ttkbootstrap as tb
from ttkbootstrap.constants import *
from openpyxl import load_workbook

# =====================================================
# Rutas fijas (plantillas & soffice)
# =====================================================
APP_OWNER = "Angel"
from pathlib import Path
import sys

BASE_DIR = Path(__file__).resolve().parent

PLANTILLAS_DIR = BASE_DIR / "plantillas"
if not PLANTILLAS_DIR.exists():
    PLANTILLAS_DIR = BASE_DIR  # fallback

EXCEL_LIQUIDACION = str((PLANTILLAS_DIR / "REPORTE DE LIQUIDACION.xlsx").resolve())
EXCEL_REGALIAS    = str((PLANTILLAS_DIR / "REGALIAS.xlsx").resolve())

# Mac LibreOffice
SOFFICE_PATH = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


# =====================================================
# Estado global
# =====================================================
entrega_folder = ""
items = []           
procesados = set()   

combo_mes = None

# =====================================================
# Utilidades
# =====================================================
def _parse_decimal_flexible(value):
    if value is None:
        return None
    s = str(value).upper().replace("GR", "").replace(" ", "").strip()
    if not s:
        return None
    s = re.sub(r"[^0-9,\.\-]", "", s)
    if s in ("", "-", ",", "."):
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None

def num_from_token(token: str):
    return _parse_decimal_flexible(token)

def num_from_text(txt):
    return _parse_decimal_flexible(txt)

def export_to_pdf(xlsx_path: str, outdir: str) -> str | None:
    os.makedirs(outdir, exist_ok=True)
    cmd = f'"{SOFFICE_PATH}" --headless --convert-to pdf --outdir "{outdir}" "{xlsx_path}"'
    try:
        subprocess.run(cmd, shell=True, check=True)
    except subprocess.CalledProcessError:
        return None
    base_pdf = os.path.join(outdir, os.path.splitext(os.path.basename(xlsx_path))[0] + ".pdf")
    return base_pdf if os.path.exists(base_pdf) else None

# =====================================================
# Parser del PDF de LEYES
# =====================================================
def parse_leyes_pdf(pdf_path: str) -> dict | None:
    try:
        with pdfplumber.open(pdf_path) as pdf:
            page = pdf.pages[0]
            lines = (page.extract_text() or "").splitlines()
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer el PDF:\n{pdf_path}\n\n{e}")
        return None

    sociedad = nit = barra = None
    peso_ini = peso_fin = ley_au = ley_ag = None

    for ln in lines:
        ln_u = ln.upper()
        if "PROVEEDOR:" in ln_u:
            sociedad = ln.split(":", 1)[1].strip()
        elif "NIT:" in ln_u:
            nit = ln.split(":", 1)[1].strip()
        elif "DOC:" in ln_u:
            barra = ln.split(":", 1)[1].strip()
        elif "BARRA" in ln_u and ":" not in ln_u:
            m = re.search(r"BARRA\s+([A-Z0-9\-]+)", ln_u)
            if m:
                barra = m.group(1)

    for ln in lines:
        tokens = re.findall(r"([0-9][0-9\.,]*\s*GR)", ln.upper())
        if len(tokens) >= 4:
            vals = [num_from_token(t) for t in tokens[:4]]
            if all(v is not None for v in vals):
                peso_ini, peso_fin, ley_au, ley_ag = vals
                break

    if not (sociedad and nit and barra and peso_ini and peso_fin and ley_au and ley_ag):
        return None

    return {
        "pdf_path": pdf_path,
        "sociedad": sociedad,
        "nit": nit,
        "barra": barra,
        "peso_ini": peso_ini,
        "peso_fin": peso_fin,
        "ley_au": ley_au,
        "ley_ag": ley_ag,
    }

# =====================================================
# Regalías (una hoja, meses en columna A, valores en C y D)
# =====================================================
def obtener_lista_meses():
    try:
        wb = load_workbook(EXCEL_REGALIAS, data_only=True)
        ws = wb.active
        meses = []
        for row in ws.iter_rows(min_row=2, max_col=1):
            mes = row[0].value
            if mes:
                meses.append(str(mes).strip())
        wb.close()
        return meses
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer REGALIAS.xlsx:\n{e}")
        return []

def leer_regalias_mes(mes: str):
    try:
        wb = load_workbook(EXCEL_REGALIAS, data_only=True)
        ws = wb.active
        mes_norm = mes.strip().lower()
        for row in ws.iter_rows(min_row=2):
            a_val = row[0].value
            if a_val and str(a_val).strip().lower() == mes_norm:
                val_c = row[2].value
                val_d = row[3].value
                wb.close()
                return val_c, val_d
        wb.close()
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo leer {mes} en REGALIAS.xlsx:\n{e}")
    return None, None

def escribir_regalias_mes(mes: str, val_c, val_d):
    try:
        wb = load_workbook(EXCEL_REGALIAS)
        ws = wb.active
        mes_norm = mes.strip().lower()
        row_target = None
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, 1).value and str(ws.cell(r, 1).value).strip().lower() == mes_norm:
                row_target = r
                break
        if not row_target:
            row_target = ws.max_row + 1
            ws.cell(row=row_target, column=1, value=mes)
        ws.cell(row=row_target, column=3, value=val_c)
        ws.cell(row=row_target, column=4, value=val_d)
        wb.save(EXCEL_REGALIAS)
        return True
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo guardar {mes} en REGALIAS.xlsx:\n{e}")
        return False

def aplicar_regalias():
    mes = combo_mes.get()
    if not mes:
        messagebox.showwarning("Atención", "Selecciona un mes.")
        return
    val_c, val_d = leer_regalias_mes(mes)
    if val_c is None and val_d is None:
        messagebox.showwarning("Atención", f"No se encontraron valores para {mes}.")
        return
    try:
        wb = load_workbook(EXCEL_LIQUIDACION)
        ws = wb.active
        ws["C53"] = val_c
        ws["E53"] = val_d
        wb.save(EXCEL_LIQUIDACION)
        messagebox.showinfo("Éxito", f"Aplicadas regalías de {mes}: Au={val_c}, Ag={val_d}")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudieron aplicar regalías:\n{e}")

def abrir_popup_regalias():
    mes = combo_mes.get()
    if not mes:
        messagebox.showwarning("Atención", "Selecciona un mes.")
        return
    val_c, val_d = leer_regalias_mes(mes)

    top = tb.Toplevel(root)
    top.title(f"Editar regalías - {mes}")
    top.geometry("320x200")

    tb.Label(top, text=f"Mes: {mes}", font=("Calibri", 12, "bold")).pack(pady=10)
    frame = tb.Frame(top, padding=10)
    frame.pack(fill="x")

    tb.Label(frame, text="Regalía Au (C):").grid(row=0, column=0, padx=5, pady=5, sticky="e")
    e_c = tb.Entry(frame); e_c.grid(row=0, column=1)
    if val_c is not None: e_c.insert(0, str(val_c))

    tb.Label(frame, text="Regalía Ag (D):").grid(row=1, column=0, padx=5, pady=5, sticky="e")
    e_d = tb.Entry(frame); e_d.grid(row=1, column=1)
    if val_d is not None: e_d.insert(0, str(val_d))

    def guardar():
        c_new = num_from_text(e_c.get())
        d_new = num_from_text(e_d.get())
        ok = escribir_regalias_mes(mes, c_new, d_new)
        if ok:
            messagebox.showinfo("Éxito", f"Regalías actualizadas para {mes}")
            top.destroy()

    tb.Button(top, text="💾 Guardar", bootstyle=SUCCESS, command=guardar).pack(pady=15)

# =====================================================
# Carga de PDFs desde carpeta ENTREGA
# =====================================================
def seleccionar_carpeta():
    global entrega_folder, items, procesados
    folder = filedialog.askdirectory(title="Selecciona la carpeta de ENTREGA")
    if not folder: return
    entrega_folder = folder
    procesados = set(); items = []

    candidatos = []
    leyes_dir = os.path.join(entrega_folder, "LEYES")
    if os.path.isdir(leyes_dir):
        candidatos += [os.path.join(leyes_dir, f) for f in os.listdir(leyes_dir) if f.lower().endswith(".pdf")]
    candidatos += [os.path.join(entrega_folder, f) for f in os.listdir(entrega_folder)
                   if f.lower().endswith(".pdf") and f.upper().startswith("REPORTE LEYES -")]

    for pdf in sorted(set(candidatos)):
        data = parse_leyes_pdf(pdf)
        if data: items.append(data)

    tabla_items.delete(*tabla_items.get_children())
    for it in items:
        tabla_items.insert("", "end", values=(os.path.basename(it["pdf_path"]), it["sociedad"], it["barra"]))

    tabla_proc.delete(*tabla_proc.get_children())

    if not items:
        messagebox.showwarning("Atención", "No se encontraron PDFs válidos.")

# =====================================================
# Escribir en Excel y generar PDF de Boletín (uno)
# =====================================================
def procesar_seleccion():
    if not entrega_folder:
        messagebox.showwarning("Atención", "Selecciona carpeta ENTREGA.")
        return
    sel = tabla_items.selection()
    if not sel:
        messagebox.showwarning("Atención", "Selecciona un PDF.")
        return
    idx = tabla_items.index(sel[0]); data = items[idx]
    if data["pdf_path"] in procesados:
        return

    aplicar_regalias()

    try:
        wb = load_workbook(EXCEL_LIQUIDACION)
        ws = wb.active
        ws["C16"] = data["barra"]
        ws["I12"] = data["sociedad"]
        ws["I14"] = data["nit"]
        ws["C24"] = data["peso_ini"]
        ws["G24"] = data["peso_fin"]
        ws["H32"] = data["ley_au"]
        ws["L32"] = data["ley_ag"]
        wb.save(EXCEL_LIQUIDACION)
    except Exception as e:
        messagebox.showerror("Error", f"No se pudo escribir en Liquidación:\n{e}")
        return

    boletines_dir = os.path.join(entrega_folder, "BOLETINES")
    pdf_generado = export_to_pdf(EXCEL_LIQUIDACION, boletines_dir)
    if pdf_generado:
        destino = os.path.join(boletines_dir, f"BOLETIN - {data['barra']}.pdf")
        if os.path.exists(destino): os.remove(destino)
        os.rename(pdf_generado, destino)
        procesados.add(data["pdf_path"])
        tabla_proc.insert("", "end", values=(os.path.basename(data["pdf_path"]), data["barra"], "✅"))
        messagebox.showinfo("Éxito", f"Boletín generado:\n{destino}")

# =====================================================
# Procesar TODOS
# =====================================================
def procesar_todos():
    if not entrega_folder:
        messagebox.showwarning("Atención", "Selecciona carpeta ENTREGA.")
        return
    if not items:
        messagebox.showwarning("Atención", "No hay PDFs para procesar.")
        return

    ok = 0
    aplicar_regalias()

    for data in items:
        if data["pdf_path"] in procesados: continue

        wb = load_workbook(EXCEL_LIQUIDACION)
        ws = wb.active
        ws["C16"] = data["barra"]
        ws["I12"] = data["sociedad"]
        ws["I14"] = data["nit"]
        ws["C24"] = data["peso_ini"]
        ws["G24"] = data["peso_fin"]
        ws["H32"] = data["ley_au"]
        ws["L32"] = data["ley_ag"]
        wb.save(EXCEL_LIQUIDACION)

        boletines_dir = os.path.join(entrega_folder, "BOLETINES")
        pdf_generado = export_to_pdf(EXCEL_LIQUIDACION, boletines_dir)
        if pdf_generado:
            destino = os.path.join(boletines_dir, f"BOLETIN - {data['barra']}.pdf")
            if os.path.exists(destino): os.remove(destino)
            os.rename(pdf_generado, destino)

            wb = load_workbook(EXCEL_LIQUIDACION)
            ws = wb.active
            for cell in ["C16","I12","I14","C24","G24","H32","L32"]:
                ws[cell] = ""
            wb.save(EXCEL_LIQUIDACION)

            procesados.add(data["pdf_path"])
            tabla_proc.insert("", "end", values=(os.path.basename(data["pdf_path"]), data["barra"], "✅"))
            ok += 1

    messagebox.showinfo("Resultado", f"Boletines generados: {ok}")

# =====================================================
# Parámetros manuales (C73, G73, J73)
# =====================================================
def guardar_parametros():
    try:
        wb = load_workbook(EXCEL_LIQUIDACION)
        ws = wb.active
        ws["C73"] = num_from_text(campo_dolar.get())
        ws["G73"] = num_from_text(campo_ozau.get())
        ws["J73"] = num_from_text(campo_ozag.get())
        wb.save(EXCEL_LIQUIDACION)
        messagebox.showinfo("Éxito", "Parámetros guardados en Liquidación (C73, G73, J73).")
    except Exception as e:
        messagebox.showerror("Error", f"No se pudieron guardar parámetros:\n{e}")

# =====================================================
# Interfaz (ttkbootstrap)
# =====================================================

BG = "#F4F7FA"
PANEL = "#FFFFFF"
INK = "#172033"
MUTED = "#667085"
BORDER = "#D8E0EA"
ACCENT = "#0E3662"

def configure_enterprise_style(window):
    window.configure(bg=BG)
    style = tb.Style()
    style.configure("App.TFrame", background=BG)
    style.configure("Card.TFrame", background=PANEL, relief="solid", borderwidth=1)
    style.configure("Header.TFrame", background=ACCENT)
    style.configure("Title.TLabel", background=ACCENT, foreground="white", font=("Calibri", 23, "bold"))
    style.configure("Subtitle.TLabel", background=ACCENT, foreground="#DDE7F2", font=("Calibri", 11))
    style.configure("Section.TLabel", background=PANEL, foreground=INK, font=("Calibri", 13, "bold"))
    style.configure("Muted.TLabel", background=PANEL, foreground=MUTED, font=("Calibri", 10))
    style.configure("Field.TLabel", background=PANEL, foreground=INK, font=("Calibri", 12, "bold"))
    style.configure("Treeview", rowheight=28, font=("Calibri", 11))
    style.configure("Treeview.Heading", font=("Calibri", 11, "bold"))

def make_header(parent, title, subtitle):
    header = ttk.Frame(parent, style="Header.TFrame", padding=(24, 18))
    header.pack(fill="x", pady=(0, 14))
    ttk.Label(header, text=title, style="Title.TLabel").pack(anchor="w")
    ttk.Label(header, text=subtitle, style="Subtitle.TLabel").pack(anchor="w", pady=(3, 0))
    return header

def animate_window(window, step=0):
    try:
        alpha = min(1.0, 0.88 + step * 0.025)
        window.attributes("-alpha", alpha)
        if alpha < 1.0:
            window.after(20, lambda: animate_window(window, step + 1))
    except Exception:
        pass

def wire_button_hover(button, normal, hover):
    def enter(_event=None):
        try:
            if str(button.cget("state")) != "disabled":
                button.configure(bootstyle=hover)
        except Exception:
            pass
    def leave(_event=None):
        try:
            button.configure(bootstyle=normal)
        except Exception:
            pass
    button.bind("<Enter>", enter)
    button.bind("<Leave>", leave)
    return button

root = tb.Window(themename="flatly")
root.title(f"Angel | BOLETINES")
root.geometry("1140x780")
root.minsize(1080, 740)
configure_enterprise_style(root)

main = tb.Frame(root, style="App.TFrame", padding=14)
main.pack(fill="both", expand=True)
make_header(main, "GENERAR BOLETINES", "Sistema Angel | Liquidación, cierre de venta y regalías")

frame_top = tb.Frame(main, style="Card.TFrame", padding=16)
frame_top.pack(fill="x", pady=(0, 10))
frame_top.columnconfigure(1, weight=1)
btn_select_entrega = tb.Button(frame_top, text="SELECCIONAR ENTREGA", bootstyle=PRIMARY, command=seleccionar_carpeta)
btn_select_entrega.grid(row=0, column=0, padx=(0, 14), ipady=5, sticky="ew")
wire_button_hover(btn_select_entrega, PRIMARY, SUCCESS)
tb.Label(frame_top, text="Detecta automáticamente los reportes de leyes disponibles", style="Muted.TLabel").grid(row=0, column=1, sticky="w")

body = tb.Frame(main, style="App.TFrame")
body.pack(fill="both", expand=True)
body.columnconfigure(0, weight=3)
body.columnconfigure(1, weight=2)
body.rowconfigure(0, weight=1)

left = tb.Frame(body, style="Card.TFrame", padding=16)
left.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
tb.Label(left, text="Sociedades detectadas", style="Section.TLabel").pack(anchor="w", pady=(0, 10))
cols = ("pdf", "sociedad", "barra")
tabla_items = ttk.Treeview(left, columns=cols, show="headings", height=9)
for c in cols:
    tabla_items.heading(c, text=c.capitalize())
tabla_items.column("pdf", width=270)
tabla_items.column("sociedad", width=360)
tabla_items.column("barra", width=120, anchor="center")
tabla_items.pack(fill="both", expand=True)

right = tb.Frame(body, style="App.TFrame")
right.grid(row=0, column=1, sticky="nsew")

frame_param = tb.Frame(right, style="Card.TFrame", padding=16)
frame_param.pack(fill="x", pady=(0, 10))
frame_param.columnconfigure(1, weight=1)
tb.Label(frame_param, text="Cierre de venta", style="Section.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 10))
tb.Label(frame_param, text="DÓLAR", style="Field.TLabel").grid(row=1, column=0, padx=(0, 8), pady=6, sticky="e")
campo_dolar = tb.Entry(frame_param); campo_dolar.grid(row=1, column=1, padx=5, pady=6, sticky="ew")
tb.Label(frame_param, text="OZ AU", style="Field.TLabel").grid(row=2, column=0, padx=(0, 8), pady=6, sticky="e")
campo_ozau = tb.Entry(frame_param); campo_ozau.grid(row=2, column=1, padx=5, pady=6, sticky="ew")
tb.Label(frame_param, text="OZ AG", style="Field.TLabel").grid(row=3, column=0, padx=(0, 8), pady=6, sticky="e")
campo_ozag = tb.Entry(frame_param); campo_ozag.grid(row=3, column=1, padx=5, pady=6, sticky="ew")
btn_guardar_param = tb.Button(frame_param, text="GUARDAR PARÁMETROS", bootstyle=SUCCESS, command=guardar_parametros)
btn_guardar_param.grid(row=4, column=0, columnspan=2, pady=(12, 0), ipady=5, sticky="ew")
wire_button_hover(btn_guardar_param, SUCCESS, PRIMARY)

frame_reg = tb.Frame(right, style="Card.TFrame", padding=16)
frame_reg.pack(fill="x", pady=(0, 10))
frame_reg.columnconfigure(1, weight=1)
tb.Label(frame_reg, text="Regalías", style="Section.TLabel").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 10))
tb.Label(frame_reg, text="MES", style="Field.TLabel").grid(row=1, column=0, padx=(0, 8), pady=6, sticky="e")
combo_mes = tb.Combobox(frame_reg, values=obtener_lista_meses(), state="readonly")
combo_mes.grid(row=1, column=1, columnspan=2, padx=5, pady=6, sticky="ew")
btn_aplicar_reg = tb.Button(frame_reg, text="APLICAR", bootstyle=SUCCESS, command=aplicar_regalias)
btn_aplicar_reg.grid(row=2, column=0, columnspan=2, padx=(0, 5), pady=(10, 0), ipady=5, sticky="ew")
wire_button_hover(btn_aplicar_reg, SUCCESS, PRIMARY)
btn_editar_reg = tb.Button(frame_reg, text="EDITAR", bootstyle=PRIMARY, command=abrir_popup_regalias)
btn_editar_reg.grid(row=2, column=2, padx=(5, 0), pady=(10, 0), ipady=5, sticky="ew")
wire_button_hover(btn_editar_reg, PRIMARY, INFO)

frame_btn = tb.Frame(right, style="Card.TFrame", padding=16)
frame_btn.pack(fill="x", pady=(0, 10))
tb.Label(frame_btn, text="Acciones", style="Section.TLabel").pack(anchor="w", pady=(0, 10))
btn_procesar_sel = tb.Button(frame_btn, text="GENERAR BOLETÍN SELECCIONADO", bootstyle=SUCCESS, command=procesar_seleccion)
btn_procesar_sel.pack(fill="x", ipady=6, pady=(0, 8))
wire_button_hover(btn_procesar_sel, SUCCESS, PRIMARY)
btn_procesar_todos = tb.Button(frame_btn, text="GENERAR TODOS", bootstyle=PRIMARY, command=procesar_todos)
btn_procesar_todos.pack(fill="x", ipady=6)
wire_button_hover(btn_procesar_todos, PRIMARY, SUCCESS)

frame_proc = tb.Frame(main, style="Card.TFrame", padding=16)
frame_proc.pack(fill="both", expand=True, pady=(10, 0))
tb.Label(frame_proc, text="Boletines generados", style="Section.TLabel").pack(anchor="w", pady=(0, 10))
cols2 = ("pdf", "barra", "estado")
tabla_proc = ttk.Treeview(frame_proc, columns=cols2, show="headings", height=4)
for c in cols2:
    tabla_proc.heading(c, text=c.capitalize())
tabla_proc.column("pdf", width=520)
tabla_proc.column("barra", width=150, anchor="center")
tabla_proc.column("estado", width=120, anchor="center")
tabla_proc.pack(fill="both", expand=True)

animate_window(root)
root.mainloop()
