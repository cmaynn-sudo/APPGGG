import ttkbootstrap as ttk
from ttkbootstrap.constants import *
from tkinter import messagebox
from openpyxl import load_workbook
from openpyxl.styles import Font
import os
import re
from datetime import datetime
from pathlib import Path
import sys

# =============================
# CONSTANTES BÁSICAS
# =============================
APP_OWNER = "Angel"
FILA_INICIO = 6
FILA_FIN = 13

SHEET_SOCIEDADES_PREFERRED = "SOCIEDADES"
SHEET_PRELIMINARES_PREFERRED = "PRELIMINARES"
SHEET_RECIBO_PREFERRED = "RECIBO"

SOFFICE_CMD_PRIMARY = r"C:\Program Files\LibreOffice\program\soffice.exe"
SOFFICE_CMD_FALLBACK = r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"

# =============================
# RUTAS DE ARCHIVOS (DINÁMICAS LOCAL)
# =============================
BASE_DIR = Path(__file__).resolve().parent

TEMPLATES_DIR = BASE_DIR / "plantillas"
if not TEMPLATES_DIR.exists():
    TEMPLATES_DIR = BASE_DIR

FILE_PRELIMINARES = str(TEMPLATES_DIR / "PRELIMINARES.xlsx")
FILE_RECIBO       = str(TEMPLATES_DIR / "RECIBO DE METALES.xlsx")

SOCIEDADES_NAME = "SOCIEDADES.xlsx"
LOCAL_SOCIEDADES = str((BASE_DIR / SOCIEDADES_NAME).resolve())

def verificar_plantillas() -> bool:
    errores = []
    if not os.path.exists(FILE_PRELIMINARES):
        errores.append(f"No se encontró PRELIMINARES.xlsx en:\n{FILE_PRELIMINARES}")
    if not os.path.exists(FILE_RECIBO):
        errores.append(f"No se encontró RECIBO DE METALES.xlsx en:\n{FILE_RECIBO}")
    if not os.path.exists(LOCAL_SOCIEDADES):
        errores.append(f"No se encontró SOCIEDADES.xlsx en:\n{LOCAL_SOCIEDADES}")

    if errores:
        messagebox.showerror(
            "Faltan archivos",
            "\n\n".join(errores) +
            "\n\nSolución:\n"
            "- Crea la carpeta 'plantillas' y pon ahí las plantillas, o\n"
            "- Pon las plantillas junto al .py.\n"
            "- SOCIEDADES.xlsx debe quedar junto a este .py (misma carpeta)."
        )
        return False
    return True

# =============================
# ESTADO GLOBAL
# =============================
entrega_global = None
fila_actual = FILA_INICIO
numero_entrega = None

# =============================
# HELPERS
# =============================
def open_wb_and_ws(file_path: str, preferred_sheet_name: str):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"No se encontró el archivo:\n{file_path}")
    wb = load_workbook(file_path)
    if preferred_sheet_name in wb.sheetnames:
        ws = wb[preferred_sheet_name]
    else:
        if not wb.sheetnames:
            wb.close()
            raise ValueError(f"El archivo '{os.path.basename(file_path)}' no contiene hojas.")
        ws = wb[wb.sheetnames[0]]
    return wb, ws

def _get_soffice_cmd() -> str:
    # Windows
    if os.name == "nt":
        if os.path.exists(SOFFICE_CMD_PRIMARY):
            return SOFFICE_CMD_PRIMARY
        if os.path.exists(SOFFICE_CMD_FALLBACK):
            return SOFFICE_CMD_FALLBACK
        return ""

    # macOS
    if sys.platform == "darwin":
        mac = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        return mac if os.path.exists(mac) else "soffice"

    # Linux
    return "soffice"

def export_to_pdf_with_lo(xlsx_path: str, desired_pdf_path: str) -> bool:
    outdir = os.path.dirname(desired_pdf_path)
    base_pdf_name = os.path.splitext(os.path.basename(xlsx_path))[0] + ".pdf"
    generated_pdf_path = os.path.join(outdir, base_pdf_name)

    soffice = _get_soffice_cmd()
    if not soffice:
        messagebox.showerror(
            "LibreOffice no encontrado",
            "No se encontró LibreOffice.\n\n"
            f"PRIMARY:\n{SOFFICE_CMD_PRIMARY}\n\n"
            f"FALLBACK:\n{SOFFICE_CMD_FALLBACK}"
        )
        return False

    try:
        if os.path.exists(generated_pdf_path):
            os.remove(generated_pdf_path)
    except:
        pass

    cmd = f'"{soffice}" --headless --convert-to pdf --outdir "{outdir}" "{xlsx_path}"'
    os.system(cmd)

    if os.path.exists(generated_pdf_path):
        try:
            if os.path.exists(desired_pdf_path):
                os.remove(desired_pdf_path)
            os.rename(generated_pdf_path, desired_pdf_path)
            return True
        except:
            return os.path.exists(desired_pdf_path)

    return False

def parse_decimal_input(value: str) -> float:
    s = str(value).strip()
    if not s:
        raise ValueError("valor vacio")
    s = re.sub(r"[^0-9,\.\-]", "", s)
    if s in ("", "-", ",", "."):
        raise ValueError("valor invalido")
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    return float(s)

def limpiar_interfaz_campos():
    combo_proveedor.set("")
    entry_peso_inicial.delete(0, "end")
    entry_peso_post.delete(0, "end")
    entry_muestras.delete(0, "end")

def obtener_sociedades():
    try:
        wb, ws = open_wb_and_ws(LOCAL_SOCIEDADES, SHEET_SOCIEDADES_PREFERRED)
        sociedades = []
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            if row and row[0]:
                sociedades.append(str(row[0]).strip())
        wb.close()
        return sociedades
    except Exception as e:
        messagebox.showerror("Error", f"No se pudieron cargar las sociedades (LOCAL):\n{e}")
        return []

def obtener_datos_sociedad(nombre_sociedad: str):
    wb, ws = open_wb_and_ws(LOCAL_SOCIEDADES, SHEET_SOCIEDADES_PREFERRED)

    prefijo = nit = rucom = municipio = ""
    consecutivo = 0
    row_index = None

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
        if row and row[0].value and str(row[0].value).strip() == nombre_sociedad:
            prefijo     = str(row[4].value).strip() if row[4].value else ""
            consecutivo = int(row[5].value) if row[5].value else 0
            nit         = str(row[1].value).strip() if row[1].value else ""
            rucom       = str(row[2].value).strip() if row[2].value else ""
            municipio   = str(row[3].value).strip() if row[3].value else ""
            row_index   = idx
            break

    wb.close()
    return prefijo, consecutivo, nit, rucom, municipio, row_index

def actualizar_consecutivo_sociedad(row_index: int, nuevo_valor: int):
    wb, ws = open_wb_and_ws(LOCAL_SOCIEDADES, SHEET_SOCIEDADES_PREFERRED)
    ws[f"F{row_index}"] = nuevo_valor
    wb.save(LOCAL_SOCIEDADES)
    wb.close()

# =============================
# ACCIONES UI
# =============================
def guardar_entrega():
    global entrega_global, numero_entrega, fila_actual

    numero_entrega = entry_entrega.get().strip()
    if not numero_entrega:
        messagebox.showwarning("Campo vacío", "Debe ingresar un valor para la ENTREGA")
        return

    fila_actual = FILA_INICIO

    # --- NUEVA RUTA AUTOMÁTICA (AÑO/MES) EN DOCUMENTS ---
    hoy = datetime.now()
    anio = str(hoy.year)

    MESES = {
        "January": "ENERO", "February": "FEBRERO", "March": "MARZO",
        "April": "ABRIL", "May": "MAYO", "June": "JUNIO",
        "July": "JULIO", "August": "AGOSTO", "September": "SEPTIEMBRE",
        "October": "OCTUBRE", "November": "NOVIEMBRE", "December": "DICIEMBRE"
    }
    mes_en = hoy.strftime("%B")
    mes = MESES.get(mes_en, mes_en).upper()

    # Base: ~/Documents/CI GREEN GLOBAL/{AÑO}/{MES}/
    base_mes = Path.home() / "Documents" / "CI GREEN GLOBAL" / anio / mes
    base_mes.mkdir(parents=True, exist_ok=True)

    # ✅ ENTREGA con número + fecha (formato que quieres)
    fecha_hoy = hoy.strftime("%Y-%m-%d")
    nombre_carpeta = f"ENTREGA °{numero_entrega} - ({fecha_hoy})"
    carpeta = base_mes / nombre_carpeta
    carpeta.mkdir(exist_ok=True)

    entrega_global = str(carpeta)
    # ----------------------------------------------------

    messagebox.showinfo("Entrega creada", f"Entrega creada localmente en:\n{entrega_global}")

    combo_proveedor.config(state="readonly")
    entry_peso_inicial.config(state="normal")
    entry_peso_post.config(state="normal")
    entry_muestras.config(state="normal")
    btn_agregar.config(state="normal")
    btn_guardar_excel.config(state="normal")
    btn_generar.config(state="normal")

def guardar_en_excel():
    global fila_actual, entrega_global

    if not entrega_global:
        messagebox.showwarning("Sin entrega", "Debe guardar primero el número de ENTREGA")
        return

    proveedor    = combo_proveedor.get().strip()
    peso_inicial = entry_peso_inicial.get().strip()
    peso_post    = entry_peso_post.get().strip()
    muestras     = entry_muestras.get().strip()

    if proveedor in ("", "Cargando..."):
        messagebox.showwarning("Proveedor", "Seleccione un proveedor válido.")
        return

    if not peso_inicial or not peso_post:
        messagebox.showwarning("Campos vacíos", "Complete Peso inicial y Peso post fundición antes de guardar")
        return

    try:
        peso_inicial_f = parse_decimal_input(peso_inicial)
        peso_post_f    = parse_decimal_input(peso_post)
        muestras_f     = parse_decimal_input(muestras) if muestras else 0.0
    except ValueError:
        messagebox.showerror("Error", "Los valores de pesos y muestras deben ser numéricos")
        return

    if fila_actual > FILA_FIN:
        messagebox.showwarning("Límite alcanzado", "Solo se pueden guardar hasta 8 sociedades (filas 6 a 13).")
        return

    try:
        prefijo, consecutivo, nit, rucom, municipio, row_index = obtener_datos_sociedad(proveedor)
        if not prefijo or row_index is None:
            messagebox.showerror("Error", f"No se encontró prefijo/registro para {proveedor}")
            return

        codigo = f"{prefijo}-{consecutivo+1}"

        # --------- PRELIMINARES ----------
        wb_p, ws_p = open_wb_and_ws(FILE_PRELIMINARES, SHEET_PRELIMINARES_PREFERRED)
        ws_p[f"B{fila_actual}"] = proveedor
        ws_p[f"C{fila_actual}"] = codigo
        ws_p[f"D{fila_actual}"] = peso_inicial_f
        ws_p[f"E{fila_actual}"] = peso_post_f
        ws_p[f"H{fila_actual}"] = muestras_f
        for col in ["B", "C", "D", "E", "H"]:
            ws_p[f"{col}{fila_actual}"].font = Font(name="Calibri", size=11)
        wb_p.save(FILE_PRELIMINARES)
        wb_p.close()

        # --------- RECIBO DE METALES ----------
        wb_r, ws_r = open_wb_and_ws(FILE_RECIBO, SHEET_RECIBO_PREFERRED)
        ws_r["I7"]  = proveedor
        ws_r["I9"]  = nit
        ws_r["I11"] = rucom
        ws_r["I13"] = municipio
        ws_r["A25"] = peso_inicial_f
        ws_r["C25"] = peso_post_f
        ws_r["B16"] = codigo
        for c in ["I7", "I9", "I11", "I13", "A25", "C25", "B16"]:
            ws_r[c].font = Font(name="Calibri", size=11, bold=True)
        wb_r.save(FILE_RECIBO)
        wb_r.close()

        # --------- GENERAR PDF RECIBO ----------
        nombre_pdf  = f"{proveedor} ({codigo}).pdf"
        ruta_pdf    = os.path.join(entrega_global, nombre_pdf)

        ok_pdf = export_to_pdf_with_lo(FILE_RECIBO, ruta_pdf)
        if not ok_pdf:
            messagebox.showerror("Error", f"No se pudo generar el PDF de RECIBO para {proveedor}.")
        else:
            # Limpia plantilla RECIBO
            wb_r2, ws_r2 = open_wb_and_ws(FILE_RECIBO, SHEET_RECIBO_PREFERRED)
            for c in ["I7", "I9", "I11", "I13", "A25", "C25", "B16"]:
                ws_r2[c] = None
            wb_r2.save(FILE_RECIBO)
            wb_r2.close()

        # --------- ACTUALIZAR CONSECUTIVO LOCAL ----------
        actualizar_consecutivo_sociedad(row_index, consecutivo+1)

        fila_actual += 1
        limpiar_interfaz_campos()

        messagebox.showinfo("Éxito", f"Guardado en PRELIMINARES y PDF RECIBO generado:\n{ruta_pdf}")

    except Exception as e:
        messagebox.showerror("Error", f"Ocurrió un error al guardar:\n{e}")

def generar_pdf_preliminares():
    global fila_actual, entrega_global, numero_entrega

    if not entrega_global:
        messagebox.showwarning("Sin entrega", "Debe ingresar la ENTREGA antes de generar el PDF.")
        return

    try:
        fecha_hoy = datetime.now().strftime("%Y-%m-%d")
        nombre_pdf = f"PRELIMINARES - {numero_entrega} ({fecha_hoy}).pdf"
        ruta_pdf   = os.path.join(entrega_global, nombre_pdf)

        ok_pdf = export_to_pdf_with_lo(FILE_PRELIMINARES, ruta_pdf)
        if not ok_pdf:
            messagebox.showerror("Error", "No se pudo generar el PDF de PRELIMINARES.\nVerifique LibreOffice.")
            return

        # LIMPIAR plantilla PRELIMINARES
        wb_p, ws_p = open_wb_and_ws(FILE_PRELIMINARES, SHEET_PRELIMINARES_PREFERRED)
        for f in range(FILA_INICIO, FILA_FIN + 1):
            for col in ["B", "C", "D", "E", "H"]:
                ws_p[f"{col}{f}"] = None
        wb_p.save(FILE_PRELIMINARES)
        wb_p.close()

        # Reset estado UI
        fila_actual = FILA_INICIO
        entrega_global = None
        entry_entrega.delete(0, "end")
        limpiar_interfaz_campos()

        combo_proveedor.config(state="disabled")
        entry_peso_inicial.config(state="disabled")
        entry_peso_post.config(state="disabled")
        entry_muestras.config(state="disabled")
        btn_agregar.config(state="disabled")
        btn_guardar_excel.config(state="disabled")
        btn_generar.config(state="disabled")

        messagebox.showinfo("Éxito", f"PDF PRELIMINARES generado en:\n{ruta_pdf}\n\nPlantilla limpia.")

    except Exception as e:
        messagebox.showerror("Error", f"Ocurrió un error al generar el PDF de PRELIMINARES:\n{e}")

# =============================
# INTERFAZ ttkbootstrap
# =============================

BG = "#F4F7FA"
PANEL = "#FFFFFF"
INK = "#172033"
MUTED = "#667085"
BORDER = "#D8E0EA"
ACCENT = "#0E3662"

def configure_enterprise_style(window):
    window.configure(bg=BG)
    style = ttk.Style()
    style.configure("App.TFrame", background=BG)
    style.configure("Card.TFrame", background=PANEL, relief="solid", borderwidth=1)
    style.configure("Header.TFrame", background=ACCENT)
    style.configure("Title.TLabel", background=ACCENT, foreground="white", font=("Calibri", 23, "bold"))
    style.configure("Subtitle.TLabel", background=ACCENT, foreground="#DDE7F2", font=("Calibri", 11))
    style.configure("Section.TLabel", background=PANEL, foreground=INK, font=("Calibri", 13, "bold"))
    style.configure("Muted.TLabel", background=PANEL, foreground=MUTED, font=("Calibri", 10))
    style.configure("Field.TLabel", background=PANEL, foreground=INK, font=("Calibri", 12, "bold"))
    style.configure("Treeview", rowheight=28, font=("Calibri", 11))
    style.configure("Treeview.Heading", font=("Calibri", 11, "bold"))

def make_header(parent, title, subtitle):
    header = ttk.Frame(parent, style="Header.TFrame", padding=(24, 18))
    header.pack(fill="x", pady=(0, 14))
    ttk.Label(header, text=title, style="Title.TLabel").pack(anchor="w")
    ttk.Label(header, text=subtitle, style="Subtitle.TLabel").pack(anchor="w", pady=(3, 0))
    return header

def animate_window(window, step=0):
    try:
        alpha = min(1.0, 0.88 + step * 0.025)
        window.attributes("-alpha", alpha)
        if alpha < 1.0:
            window.after(20, lambda: animate_window(window, step + 1))
    except Exception:
        pass

def wire_button_hover(button, normal, hover):
    def enter(_event=None):
        try:
            if str(button.cget("state")) != "disabled":
                button.configure(bootstyle=hover)
        except Exception:
            pass
    def leave(_event=None):
        try:
            button.configure(bootstyle=normal)
        except Exception:
            pass
    button.bind("<Enter>", enter)
    button.bind("<Leave>", leave)
    return button

app = ttk.Window(themename="flatly")
app.title(f"Angel | PRELIMINARES + RECIBO")
app.geometry("940x600")
app.minsize(900, 580)
configure_enterprise_style(app)

if not verificar_plantillas():
    app.destroy()
    raise SystemExit(0)

main = ttk.Frame(app, style="App.TFrame", padding=18)
main.pack(fill="both", expand=True)

make_header(main, "PRELIMINARES Y RECIBO", "Sistema Angel | Captura de entrega, sociedades y pesos con punto o coma decimal")

content = ttk.Frame(main, style="Card.TFrame", padding=20)
content.pack(fill="x")
content.columnconfigure(1, weight=1)
content.columnconfigure(2, minsize=230)

ttk.Label(content, text="Datos de la entrega", style="Section.TLabel").grid(row=0, column=0, columnspan=3, sticky="w", pady=(0, 16))

ttk.Label(content, text="Número de ENTREGA", style="Field.TLabel").grid(row=1, column=0, padx=(0, 14), pady=9, sticky="e")
entry_entrega = ttk.Entry(content, font=("Calibri", 13))
entry_entrega.grid(row=1, column=1, padx=5, pady=9, sticky="ew")
btn_entrega = ttk.Button(content, text="GUARDAR ENTREGA", command=guardar_entrega, bootstyle=SUCCESS, width=24)
btn_entrega.grid(row=1, column=2, padx=(14, 0), pady=9, sticky="ew")
wire_button_hover(btn_entrega, SUCCESS, PRIMARY)

ttk.Label(content, text="Proveedor", style="Field.TLabel").grid(row=2, column=0, padx=(0, 14), pady=9, sticky="e")
combo_proveedor = ttk.Combobox(content, values=["Cargando..."], width=42, font=("Calibri", 13), state="disabled")
combo_proveedor.grid(row=2, column=1, columnspan=2, padx=5, pady=9, sticky="ew")

def cargar_sociedades_en_combo():
    sociedades = obtener_sociedades()
    combo_proveedor["values"] = sociedades
    combo_proveedor.set("")

app.after(200, cargar_sociedades_en_combo)

ttk.Label(content, text="Peso inicial (grs)", style="Field.TLabel").grid(row=3, column=0, padx=(0, 14), pady=9, sticky="e")
entry_peso_inicial = ttk.Entry(content, font=("Calibri", 13), state="disabled")
entry_peso_inicial.grid(row=3, column=1, padx=5, pady=9, sticky="ew")

ttk.Label(content, text="Peso post fundición (grs)", style="Field.TLabel").grid(row=4, column=0, padx=(0, 14), pady=9, sticky="e")
entry_peso_post = ttk.Entry(content, font=("Calibri", 13), state="disabled")
entry_peso_post.grid(row=4, column=1, padx=5, pady=9, sticky="ew")

ttk.Label(content, text="Muestras", style="Field.TLabel").grid(row=5, column=0, padx=(0, 14), pady=9, sticky="e")
entry_muestras = ttk.Entry(content, font=("Calibri", 13), state="disabled")
entry_muestras.grid(row=5, column=1, padx=5, pady=9, sticky="ew")
ttk.Label(content, text="Ejemplo: 12.5 o 12,5", style="Muted.TLabel").grid(row=5, column=2, padx=(14, 0), pady=9, sticky="w")

actions_card = ttk.Frame(main, style="Card.TFrame", padding=16)
actions_card.pack(fill="x", pady=(14, 0))
actions_card.columnconfigure(0, weight=1)
actions_card.columnconfigure(1, weight=1)
actions_card.columnconfigure(2, weight=1)
btn_agregar = ttk.Button(actions_card, text="AGREGAR SOCIEDAD", command=guardar_en_excel, bootstyle=INFO, state="disabled")
btn_agregar.grid(row=0, column=0, padx=(0, 8), ipady=5, sticky="ew")
wire_button_hover(btn_agregar, INFO, PRIMARY)
btn_guardar_excel = ttk.Button(actions_card, text="GUARDAR", command=guardar_en_excel, bootstyle=PRIMARY, state="disabled")
btn_guardar_excel.grid(row=0, column=1, padx=8, ipady=5, sticky="ew")
wire_button_hover(btn_guardar_excel, PRIMARY, INFO)
btn_generar = ttk.Button(actions_card, text="GENERAR PDF PRELIMINARES", command=generar_pdf_preliminares, bootstyle=SUCCESS, state="disabled")
btn_generar.grid(row=0, column=2, padx=(8, 0), ipady=5, sticky="ew")
wire_button_hover(btn_generar, SUCCESS, PRIMARY)

animate_window(app)
app.mainloop()
