from __future__ import annotations

import html
import re
import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from template_catalog import GENERATED_ASSET_DIR, GENERATED_TEMPLATE_DIR, PROJECT_DIR, TEMPLATES, TemplateSpec


EMU_PER_PIXEL = 9525
DEFAULT_ROW_HEIGHT_PT = 15


def excel_width_to_px(width: float | None) -> int:
    if width is None:
        width = 8.43
    return max(18, int(width * 7 + 5))


def row_height_to_px(height: float | None) -> int:
    if height is None:
        height = DEFAULT_ROW_HEIGHT_PT
    return max(16, int(float(height) * 1.333))


def color_to_css(color: Any) -> str | None:
    if not color:
        return None
    rgb = getattr(color, "rgb", None)
    if rgb and isinstance(rgb, str):
        rgb = rgb.upper()
        if rgb in {"00000000", "FFFFFFFF"} and getattr(color, "type", None) == "theme":
            return None
        if len(rgb) == 8:
            rgb = rgb[2:]
        if len(rgb) == 6 and rgb != "000000":
            return f"#{rgb}"
        if len(rgb) == 6 and getattr(color, "type", None) == "rgb":
            return f"#{rgb}"
    return None


def border_to_css(side: Any) -> str | None:
    style = getattr(side, "style", None)
    if not style:
        return None
    width = "2px" if style in {"medium", "thick", "double"} else "1px"
    color = color_to_css(getattr(side, "color", None)) or "#1f2937"
    line_style = "double" if style == "double" else "solid"
    return f"{width} {line_style} {color}"


def cell_style_to_css(cell: Any) -> str:
    css: list[str] = [
        "box-sizing:border-box",
        "padding:2px 4px",
    ]

    font = cell.font
    if font.name:
        css.append(f"font-family:{css_string(font.name)}")
    if font.sz:
        css.append(f"font-size:{float(font.sz) * 1.333:.1f}px")
    if font.bold:
        css.append("font-weight:700")
    if font.italic:
        css.append("font-style:italic")
    font_color = color_to_css(font.color)
    if font_color:
        css.append(f"color:{font_color}")

    fill = cell.fill
    fill_color = color_to_css(fill.fgColor)
    if getattr(fill, "fill_type", None) and fill_color:
        css.append(f"background:{fill_color}")

    alignment = cell.alignment
    if alignment.horizontal:
        css.append(f"text-align:{alignment.horizontal}")
    if alignment.vertical:
        vertical = "middle" if alignment.vertical == "center" else alignment.vertical
        css.append(f"vertical-align:{vertical}")
    if alignment.wrap_text:
        css.append("white-space:normal")
    else:
        css.append("white-space:pre-wrap")

    border = cell.border
    for attr, css_name in [
        ("left", "border-left"),
        ("right", "border-right"),
        ("top", "border-top"),
        ("bottom", "border-bottom"),
    ]:
        value = border_to_css(getattr(border, attr))
        if value:
            css.append(f"{css_name}:{value}")

    return ";".join(css)


def css_string(value: str) -> str:
    clean = value.replace("\\", "\\\\").replace('"', '\\"')
    return f'"{clean}"'


def parse_print_area(ws: Any) -> tuple[int, int, int, int]:
    print_area = ws.print_area
    if print_area:
        if isinstance(print_area, str):
            area = print_area.split("!", 1)[-1].replace("$", "").replace("'", "")
        else:
            area = str(print_area).split("!", 1)[-1].replace("$", "").replace("'", "")
        try:
            return range_boundaries(area)
        except ValueError:
            pass
    return 1, 1, ws.max_column, ws.max_row


def image_extension(image: Any) -> str:
    fmt = (getattr(image, "format", None) or "png").lower()
    if fmt == "jpeg":
        return "jpg"
    return fmt


def extract_images(ws: Any, spec: TemplateSpec, min_col: int, min_row: int) -> list[str]:
    GENERATED_ASSET_DIR.mkdir(parents=True, exist_ok=True)
    snippets: list[str] = []
    max_anchor_col = ws.max_column
    max_anchor_row = ws.max_row
    for image in getattr(ws, "_images", []):
        anchor = getattr(image, "anchor", None)
        for marker in (getattr(anchor, "_from", None), getattr(anchor, "to", None)):
            if marker:
                max_anchor_col = max(max_anchor_col, int(marker.col) + 1)
                max_anchor_row = max(max_anchor_row, int(marker.row) + 1)

    col_offsets = {min_col: 0}
    running = 0
    for col in range(min_col, max_anchor_col + 2):
        col_offsets[col] = running
        running += excel_width_to_px(ws.column_dimensions[get_column_letter(col)].width)

    row_offsets = {min_row: 0}
    running_y = 0
    for row in range(min_row, max_anchor_row + 2):
        row_offsets[row] = running_y
        running_y += row_height_to_px(ws.row_dimensions[row].height)

    for idx, image in enumerate(getattr(ws, "_images", []), start=1):
        ext = image_extension(image)
        name = f"{spec.slug}-image-{idx}.{ext}"
        out_path = GENERATED_ASSET_DIR / name
        try:
            out_path.write_bytes(image._data())
        except Exception:
            continue

        anchor = getattr(image, "anchor", None)
        marker = getattr(anchor, "_from", None)
        if not marker:
            continue
        col = int(marker.col) + 1
        row = int(marker.row) + 1
        col_offset = int(getattr(marker, "colOff", 0) or 0) / EMU_PER_PIXEL
        row_offset = int(getattr(marker, "rowOff", 0) or 0) / EMU_PER_PIXEL
        x = col_offsets.get(col, 0) + col_offset
        y = row_offsets.get(row, 0) + row_offset
        end_marker = getattr(anchor, "to", None)
        ext = getattr(anchor, "ext", None)
        if end_marker:
            end_col = int(end_marker.col) + 1
            end_row = int(end_marker.row) + 1
            end_x = col_offsets.get(end_col, x) + (int(getattr(end_marker, "colOff", 0) or 0) / EMU_PER_PIXEL)
            end_y = row_offsets.get(end_row, y) + (int(getattr(end_marker, "rowOff", 0) or 0) / EMU_PER_PIXEL)
            width = max(1, int(end_x - x))
            height = max(1, int(end_y - y))
        elif ext:
            width = max(1, int(ext.width / EMU_PER_PIXEL))
            height = max(1, int(ext.height / EMU_PER_PIXEL))
        else:
            width = int(getattr(image, "width", 120) or 120)
            height = int(getattr(image, "height", 80) or 80)
        snippets.append(
            f'<img class="excel-image" src="/static/generated/{html.escape(name)}" '
            f'style="left:{x:.0f}px;top:{y:.0f}px;width:{width}px;height:{height}px" alt="">'
        )
    return snippets


def build_merge_maps(ws: Any, bounds: tuple[int, int, int, int]) -> tuple[dict[str, tuple[int, int]], set[str]]:
    min_col, min_row, max_col, max_row = bounds
    spans: dict[str, tuple[int, int]] = {}
    covered: set[str] = set()
    for merged in ws.merged_cells.ranges:
        if merged.max_row < min_row or merged.min_row > max_row or merged.max_col < min_col or merged.min_col > max_col:
            continue
        top_left = f"{get_column_letter(merged.min_col)}{merged.min_row}"
        spans[top_left] = (merged.max_row - merged.min_row + 1, merged.max_col - merged.min_col + 1)
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                coord = f"{get_column_letter(col)}{row}"
                if coord != top_left:
                    covered.add(coord)
    return spans, covered


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str) and value.startswith("="):
        return ""
    return str(value)


def cell_html(cell: Any, data_cell: Any, fields: dict[str, str]) -> str:
    coord = cell.coordinate
    if coord in fields:
        return "{{ " + fields[coord] + " }}"
    value = display_value(data_cell.value if cell.value and str(cell.value).startswith("=") else cell.value)
    escaped = html.escape(value)
    return escaped.replace("\n", "<br>")


def render_template(spec: TemplateSpec) -> str:
    wb = load_workbook(spec.workbook)
    wb_values = load_workbook(spec.workbook, data_only=True)
    ws = wb[spec.preferred_sheet] if spec.preferred_sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    ws_values = wb_values[ws.title] if ws.title in wb_values.sheetnames else wb_values[wb_values.sheetnames[0]]

    bounds = parse_print_area(ws)
    min_col, min_row, max_col, max_row = bounds
    spans, covered = build_merge_maps(ws, bounds)
    col_widths = [excel_width_to_px(ws.column_dimensions[get_column_letter(col)].width) for col in range(min_col, max_col + 1)]
    sheet_width = sum(col_widths)
    images = extract_images(ws, spec, min_col, min_row)

    lines = [
        f'<section class="print-document" data-template="{html.escape(spec.slug)}">',
        f'<div class="doc-sheet" style="width:{sheet_width}px">',
        *images,
        '<table class="excel-sheet" aria-label="' + html.escape(spec.title) + '">',
        "<colgroup>",
    ]
    for width in col_widths:
        lines.append(f'<col style="width:{width}px">')
    lines.append("</colgroup>")

    for row in range(min_row, max_row + 1):
        height = row_height_to_px(ws.row_dimensions[row].height)
        lines.append(f'<tr style="height:{height}px">')
        for col in range(min_col, max_col + 1):
            coord = f"{get_column_letter(col)}{row}"
            if coord in covered:
                continue
            cell = ws[coord]
            data_cell = ws_values[coord]
            span = spans.get(coord)
            attrs = []
            classes = []
            if coord in spec.fields:
                classes.append("excel-field")
            if classes:
                attrs.append(f'class="{" ".join(classes)}"')
            if span:
                rowspan, colspan = span
                attrs.append(f'rowspan="{rowspan}"')
                attrs.append(f'colspan="{colspan}"')
            style = cell_style_to_css(cell)
            attrs.append(f'style="{html.escape(style, quote=True)}"')
            content = cell_html(cell, data_cell, spec.fields)
            lines.append(f"<td {' '.join(attrs)}>{content}</td>")
        lines.append("</tr>")

    lines.extend(["</table>", "</div>", "</section>"])
    wb.close()
    wb_values.close()
    return "\n".join(lines)


def export_all_templates() -> list[Path]:
    GENERATED_TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    GENERATED_ASSET_DIR.mkdir(parents=True, exist_ok=True)
    outputs: list[Path] = []
    for spec in TEMPLATES:
        if not spec.workbook.exists():
            continue
        out_path = GENERATED_TEMPLATE_DIR / f"{spec.slug}.html"
        out_path.write_text(render_template(spec), encoding="utf-8")
        outputs.append(out_path)

    for asset in ["logo.png", "firma.png"]:
        src = PROJECT_DIR / asset
        if src.exists():
            shutil.copy2(src, GENERATED_ASSET_DIR / asset)
    return outputs


if __name__ == "__main__":
    paths = export_all_templates()
    for path in paths:
        print(path.relative_to(PROJECT_DIR))
