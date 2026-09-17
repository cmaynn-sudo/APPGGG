from __future__ import annotations

import calendar
import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from template_catalog import PROJECT_DIR


ROOT_BASE = Path.home() / "Documents" / "CI GREEN GLOBAL"
SOCIEDADES_FILE = PROJECT_DIR / "SOCIEDADES.xlsx"
REGALIAS_FILE = PROJECT_DIR / "REGALIAS.xlsx"

MESES_ES = {
    1: "ENERO",
    2: "FEBRERO",
    3: "MARZO",
    4: "ABRIL",
    5: "MAYO",
    6: "JUNIO",
    7: "JULIO",
    8: "AGOSTO",
    9: "SEPTIEMBRE",
    10: "OCTUBRE",
    11: "NOVIEMBRE",
    12: "DICIEMBRE",
}
MESES_ORDEN = list(MESES_ES.values())


def parse_decimal_input(value: str) -> float:
    s = str(value).strip()
    if not s:
        raise ValueError("valor vacío")
    s = re.sub(r"[^0-9,\.\-]", "", s)
    if s in ("", "-", ",", "."):
        raise ValueError("valor inválido")

    sign = ""
    if s.startswith("-"):
        sign = "-"
        s = s[1:]
    if "-" in s:
        raise ValueError("valor inválido")

    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        separator = "," if "," in s else "." if "." in s else ""
        if separator:
            whole, fractional = s.rsplit(separator, 1)
            whole_groups = whole.split(separator)
            valid_grouping = (
                all(group.isdigit() for group in whole_groups + [fractional])
                and 1 <= len(whole_groups[0]) <= 3
                and all(len(group) == 3 for group in whole_groups[1:])
            )
            single_group_thousands = len(whole_groups) == 1 and len(whole_groups[0]) <= 3
            multi_group_thousands = len(whole_groups) > 1
            looks_like_thousands = (
                len(fractional) == 3
                and valid_grouping
                and (single_group_thousands or multi_group_thousands)
            )
            if looks_like_thousands:
                s = s.replace(separator, "")
            else:
                s = s.replace(separator, ".")
    return float(sign + s)


def mes_actual_es(today: date | None = None) -> str:
    current = today or date.today()
    return MESES_ES[current.month]


def read_sociedades() -> list[dict[str, Any]]:
    if not SOCIEDADES_FILE.exists():
        return []
    wb = load_workbook(SOCIEDADES_FILE, data_only=True, read_only=True)
    ws = wb["SOCIEDADES"] if "SOCIEDADES" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rows.append(
            {
                "sociedad": str(row[0]).strip(),
                "nit": str(row[1]).strip() if row[1] else "",
                "rucom": str(int(row[2])) if isinstance(row[2], float) else (str(row[2]).strip() if row[2] else ""),
                "municipio": str(row[3]).strip() if row[3] else "",
                "prefijo": str(row[4]).strip() if row[4] else "",
                "consecutivo": int(row[5] or 0),
            }
        )
    wb.close()
    return rows


def read_regalias() -> list[dict[str, Any]]:
    if not REGALIAS_FILE.exists():
        return []
    wb = load_workbook(REGALIAS_FILE, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        rows.append(
            {
                "mes": str(row[0]).strip(),
                "au": row[2] if len(row) > 2 else None,
                "ag": row[3] if len(row) > 3 else None,
            }
        )
    wb.close()
    return rows


def list_years() -> list[str]:
    if not ROOT_BASE.exists():
        return []
    years = [p.name for p in ROOT_BASE.iterdir() if p.is_dir() and re.match(r"^\d{4}$", p.name)]
    return sorted(years)


def list_months_for_year(year: str) -> list[str]:
    year_dir = ROOT_BASE / year
    if not year_dir.exists():
        return []
    months = []
    for p in year_dir.iterdir():
        if p.is_dir() and p.name.upper() in MESES_ORDEN:
            months.append(p.name.upper())
    return sorted(months, key=lambda m: MESES_ORDEN.index(m))


def list_entregas(year: str, month: str) -> list[dict[str, Any]]:
    base_month = ROOT_BASE / year / month
    if not base_month.exists():
        return []
    entregas = []
    for entrega_dir in sorted(base_month.iterdir(), key=lambda p: p.name):
        if not entrega_dir.is_dir() or not entrega_dir.name.upper().startswith("ENTREGA"):
            continue
        leyes_count = count_pdfs(entrega_dir / "LEYES", "REPORTE LEYES")
        boletines_count = count_pdfs(entrega_dir / "BOLETINES", "BOLETIN")
        preliminares_count = count_pdfs(entrega_dir, "PRELIMINARES")
        recibos_count = len(
            [
                f
                for f in entrega_dir.iterdir()
                if f.is_file()
                and f.suffix.lower() == ".pdf"
                and not f.name.upper().startswith(("PRELIMINARES", "REPORTE LEYES", "BOLETIN"))
            ]
        )
        entregas.append(
            {
                "name": entrega_dir.name,
                "path": str(entrega_dir),
                "preliminares": preliminares_count,
                "recibos": recibos_count,
                "leyes": leyes_count,
                "boletines": boletines_count,
            }
        )
    return entregas


def count_pdfs(folder: Path, prefix: str) -> int:
    if not folder.exists() or not folder.is_dir():
        return 0
    return len([f for f in folder.iterdir() if f.is_file() and f.suffix.lower() == ".pdf" and f.name.upper().startswith(prefix)])


def next_delivery_folder(numero_entrega: str, today: date | None = None) -> Path:
    current = today or date.today()
    year = str(current.year)
    month = mes_actual_es(current)
    folder_name = f"ENTREGA \u00b0{numero_entrega} - ({current.isoformat()})"
    return ROOT_BASE / year / month / folder_name


def json_response(data: Any) -> bytes:
    return json.dumps(data, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def current_period() -> tuple[str, str]:
    today = datetime.now().date()
    return str(today.year), MESES_ES[today.month]
